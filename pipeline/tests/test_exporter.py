"""
Unit tests for exporter module in exporter.py.
Verifies strict 6-column schema, Excel formatting, CSV export, and empty Feedback column constraint.
"""

from __future__ import annotations

import csv
import tempfile
from pathlib import Path

import openpyxl

from pipeline.exporter import (
    STRICT_OUTPUT_COLUMNS,
    export_results,
    export_to_csv,
    export_to_excel,
)


def test_strict_column_definitions():
    """Verify strict 6-column definition and exact ordering."""
    expected = (
        "Feed Title",
        "Analyst Name",
        "Classification",
        "Event Type",
        "Rationale",
        "Feedback",
    )
    assert STRICT_OUTPUT_COLUMNS == expected
    assert len(STRICT_OUTPUT_COLUMNS) == 6


def test_export_to_csv():
    """Verify CSV export produces correct header and data rows with empty feedback."""
    records = [
        {
            "Feed Title": "Factory fire at plant in Austin",
            "Analyst Name": "John Doe",
            "Classification": "Impactful",
            "Event Type": "Factory Fire",
            "Rationale": "Fire damages manufacturing assembly line.",
        },
        {
            "Feed Title": "Rosen Law Firm announces investigation",
            "Analyst Name": "Unassigned",
            "Classification": "Not Impactful",
            "Event Type": "Legal Action",
            "Rationale": "Securities shareholder solicitation ad.",
        },
    ]

    with tempfile.TemporaryDirectory() as tmp_dir:
        csv_path = Path(tmp_dir) / "output.csv"
        out = export_to_csv(records, csv_path)
        assert out.is_file()

        with open(out, mode="r", newline="", encoding="utf-8") as f:
            reader = list(csv.reader(f))

        assert len(reader) == 3  # Header + 2 data rows
        assert tuple(reader[0]) == STRICT_OUTPUT_COLUMNS

        # Row 1
        assert reader[1][0] == "Factory fire at plant in Austin"
        assert reader[1][1] == "John Doe"
        assert reader[1][2] == "Impactful"
        assert reader[1][3] == "Factory Fire"
        assert reader[1][4] == "Fire damages manufacturing assembly line."
        assert reader[1][5] == ""  # Strict empty feedback

        # Row 2
        assert reader[2][5] == ""


def test_export_to_excel():
    """Verify Excel workbook generation with styling and strict 6-column structure."""
    records = [
        {
            "Feed Title": "Strike at Port of Hamburg",
            "Analyst Name": "Garima Sinha",
            "Classification": "Impactful",
            "Event Type": "Labor Disruption",
            "Rationale": "Port strike halts container freight handling.",
        },
        {
            "Feed Title": "Routine quarterly dividend announced",
            "Analyst Name": "Unassigned",
            "Classification": "Not Impactful",
            "Event Type": "Financial Distress",
            "Rationale": "Routine financial distribution without operational disruption.",
        },
    ]

    with tempfile.TemporaryDirectory() as tmp_dir:
        xlsx_path = Path(tmp_dir) / "output.xlsx"
        out = export_to_excel(records, xlsx_path)
        assert out.is_file()

        wb = openpyxl.load_workbook(str(out), data_only=True)
        ws = wb.active
        assert ws.title == "Classified Events"

        # Check headers
        headers = [cell.value for cell in ws[1]]
        assert tuple(headers) == STRICT_OUTPUT_COLUMNS

        # Check rows count
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) == 3

        # Check Row 2 data
        row2 = rows[1]
        assert row2[0] == "Strike at Port of Hamburg"
        assert row2[1] == "Garima Sinha"
        assert row2[2] == "Impactful"
        assert row2[3] == "Labor Disruption"
        assert row2[4] == "Port strike halts container freight handling."
        assert row2[5] in ("", None)

        # Check Row 3 data
        row3 = rows[2]
        assert row3[2] == "Not Impactful"
        assert row3[5] in ("", None)

        wb.close()


def test_export_results_convenience():
    """Verify export_results creates both Excel and CSV files simultaneously."""
    records = [
        {
            "Feed Title": "Cyber attack on logistics provider",
            "Analyst Name": "Alex Chen",
            "Classification": "Impactful",
            "Event Type": "Cyber Attack",
            "Rationale": "Ransomware disrupts freight scheduling.",
        }
    ]
    with tempfile.TemporaryDirectory() as tmp_dir:
        excel_path = Path(tmp_dir) / "test_out.xlsx"
        res = export_results(records, excel_path)
        assert res["excel"].is_file()
        assert res["csv"].is_file()
        assert res["csv"].suffix == ".csv"
