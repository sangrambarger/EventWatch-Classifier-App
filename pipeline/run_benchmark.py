"""
Live Benchmark Runner for EventWatch Pipeline.
Extracts 20 rows from 04_Notified_Story_Report.xlsx and runs end-to-end live classification.
Verifies row preservation, 6-column output schema, and canonical event types.
"""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

from pipeline.config import CANONICAL_EVENT_TYPES_SET
from pipeline.exporter import STRICT_OUTPUT_COLUMNS
from pipeline.main import run_pipeline


def prepare_benchmark_file(
    source_file: Path | str = "04_Notified_Story_Report.xlsx",
    output_file: Path | str = "output/benchmark_20_input.xlsx",
    row_count: int = 20,
) -> Path:
    """Extracts first N data rows from source workbook to create a benchmark input file."""
    src = Path(source_file)
    dst = Path(output_file)
    dst.parent.mkdir(parents=True, exist_ok=True)

    wb_src = openpyxl.load_workbook(str(src), read_only=True, data_only=True)
    ws_src = wb_src.active
    assert ws_src is not None

    rows = list(ws_src.iter_rows(max_row=row_count + 1, values_only=True))
    wb_src.close()

    wb_dst = openpyxl.Workbook()
    ws_dst = wb_dst.active
    ws_dst.title = "Benchmark Events"
    for r in rows:
        ws_dst.append(list(r))
    wb_dst.save(str(dst))
    wb_dst.close()

    return dst


def verify_benchmark_output(
    excel_path: Path | str,
    csv_path: Path | str,
    expected_rows: int = 20,
) -> None:
    """Validates the output Excel and CSV against strict EventWatch schema requirements."""
    excel_p = Path(excel_path)
    csv_p = Path(csv_path)

    assert excel_p.is_file(), f"Output Excel file missing: {excel_p}"
    assert csv_p.is_file(), f"Output CSV file missing: {csv_p}"

    # Verify CSV
    with open(csv_p, mode="r", newline="", encoding="utf-8") as f:
        reader = list(csv.reader(f))

    assert len(reader) == expected_rows + 1, f"Expected {expected_rows + 1} lines in CSV, found {len(reader)}"
    header = reader[0]
    assert tuple(header) == STRICT_OUTPUT_COLUMNS, f"Header mismatch: {header} vs {STRICT_OUTPUT_COLUMNS}"

    for idx, row in enumerate(reader[1:], start=1):
        assert len(row) == 6, f"Row {idx} has {len(row)} columns, expected 6"
        feed_title, analyst, classification, event_type, rationale, feedback = row

        assert feed_title.strip(), f"Row {idx}: empty Feed Title"
        assert analyst.strip(), f"Row {idx}: empty Analyst Name"
        assert classification in ("Impactful", "Not Impactful"), f"Row {idx}: invalid classification '{classification}'"
        assert event_type in CANONICAL_EVENT_TYPES_SET, f"Row {idx}: non-canonical event type '{event_type}'"
        assert len(rationale.strip()) > 5, f"Row {idx}: rationale too short '{rationale}'"
        assert feedback == "", f"Row {idx}: Feedback column must be strictly empty, got '{feedback}'"

    # Verify Excel
    wb = openpyxl.load_workbook(str(excel_p), data_only=True)
    ws = wb.active
    excel_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    assert len(excel_rows) == expected_rows + 1
    assert tuple(excel_rows[0]) == STRICT_OUTPUT_COLUMNS
    for idx, row in enumerate(excel_rows[1:], start=1):
        assert row[2] in ("Impactful", "Not Impactful")
        assert row[3] in CANONICAL_EVENT_TYPES_SET
        assert row[5] in ("", None)


def run_benchmark() -> bool:
    """Executes the complete benchmark workflow."""
    print("\n" + "=" * 80)
    print("STARTING LIVE BENCHMARK TEST (20 ROWS FROM 04_Notified_Story_Report.xlsx)")
    print("=" * 80)

    input_path = prepare_benchmark_file(
        source_file="04_Notified_Story_Report.xlsx",
        output_file="output/benchmark_20_input.xlsx",
        row_count=20,
    )
    print(f"Benchmark input prepared: {input_path.resolve()}")

    output_excel = Path("output/benchmark_20_classified.xlsx")
    output_csv = Path("output/benchmark_20_classified.csv")

    exit_code = run_pipeline(
        input_file=input_path,
        output_file=output_excel,
        batch_size=10,
    )
    assert exit_code == 0, f"Pipeline exited with error code {exit_code}"

    verify_benchmark_output(
        excel_path=output_excel,
        csv_path=output_csv,
        expected_rows=20,
    )

    print("\n" + "=" * 80)
    print("BENCHMARK VERIFICATION PASSED: 100% SUCCESS")
    print(f"Verified 20/20 rows processed, 6 strict columns, 100% canonical taxonomy compliant.")
    print(f"Output files verified:\n  Excel: {output_excel.resolve()}\n  CSV:   {output_csv.resolve()}")
    print("=" * 80)
    return True


if __name__ == "__main__":
    run_benchmark()
