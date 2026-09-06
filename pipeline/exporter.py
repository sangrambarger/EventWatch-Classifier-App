"""
Tabular exporter module for EventWatch pipeline.
Generates production Excel (.xlsx) and CSV (.csv) reports with strict 6-column layout.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


STRICT_OUTPUT_COLUMNS: Sequence[str] = (
    "Feed Title",
    "Analyst Name",
    "Classification",
    "Event Type",
    "Rationale",
    "Feedback",
)


def export_to_csv(
    records: List[Dict[str, Any]],
    output_path: Path | str,
) -> Path:
    """Exports evaluated records to CSV with strict 6-column schema."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    with open(out, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(STRICT_OUTPUT_COLUMNS)
        for rec in records:
            writer.writerow([
                rec.get("Feed Title", ""),
                rec.get("Analyst Name", ""),
                rec.get("Classification", ""),
                rec.get("Event Type", ""),
                rec.get("Rationale", ""),
                "",  # Feedback column is explicitly empty string
            ])

    return out


def export_to_excel(
    records: List[Dict[str, Any]],
    output_path: Path | str,
) -> Path:
    """Exports evaluated records to styled Excel workbook with strict 6-column schema."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Classified Events"

    # Styling definitions
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    body_font = Font(name="Calibri", size=10)
    wrap_align = Alignment(vertical="top", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="top")

    # Classification highlight fills
    impactful_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    not_impactful_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    # Write headers
    ws.append(list(STRICT_OUTPUT_COLUMNS))
    header_row = ws[1]
    for cell in header_row:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Write data rows
    for row_idx, rec in enumerate(records, start=2):
        feed_title = str(rec.get("Feed Title", ""))
        analyst_name = str(rec.get("Analyst Name", ""))
        classification = str(rec.get("Classification", ""))
        event_type = str(rec.get("Event Type", ""))
        rationale = str(rec.get("Rationale", ""))
        feedback = ""  # Explicit empty string for every row

        row_values = [feed_title, analyst_name, classification, event_type, rationale, feedback]
        ws.append(row_values)

        current_row = ws[row_idx]
        for col_idx, cell in enumerate(current_row):
            cell.font = body_font
            cell.border = thin_border

            # Column 1 (Feed Title): wrap text
            if col_idx == 0:
                cell.alignment = wrap_align
            # Column 2 (Analyst Name): center
            elif col_idx == 1:
                cell.alignment = center_align
            # Column 3 (Classification): center + highlight fill
            elif col_idx == 2:
                cell.alignment = center_align
                if classification.strip().lower() == "impactful":
                    cell.fill = impactful_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="C00000")
                elif classification.strip().lower() == "not impactful":
                    cell.fill = not_impactful_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="385723")
            # Column 4 (Event Type): center
            elif col_idx == 3:
                cell.alignment = center_align
            # Column 5 (Rationale): wrap text
            elif col_idx == 4:
                cell.alignment = wrap_align
            # Column 6 (Feedback): empty
            elif col_idx == 5:
                cell.alignment = center_align

        ws.row_dimensions[row_idx].height = 36

    # Set column widths
    column_widths = {
        "A": 45,  # Feed Title
        "B": 18,  # Analyst Name
        "C": 16,  # Classification
        "D": 22,  # Event Type
        "E": 55,  # Rationale
        "F": 12,  # Feedback
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(str(out))
    wb.close()
    return out


def export_results(
    records: List[Dict[str, Any]],
    excel_path: Path | str,
    csv_path: Optional[Path | str] = None,
) -> Dict[str, Path]:
    """Generates both Excel and CSV output files."""
    results: Dict[str, Path] = {}
    excel_out = export_to_excel(records, excel_path)
    results["excel"] = excel_out

    if csv_path is None:
        csv_path = Path(excel_path).with_suffix(".csv")
    csv_out = export_to_csv(records, csv_path)
    results["csv"] = csv_out

    return results
